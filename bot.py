import asyncio
import logging
import os
import re
import io
from datetime import datetime, timedelta
import pytz
from aiogram import Bot, Dispatcher, F, Router, types
from aiogram.filters import Command, StateFilter
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import BotCommand, InlineKeyboardMarkup, InlineKeyboardButton, BufferedInputFile

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# DB faylidan funksiyalarni import qilish
from app.database.db import (
    init_db,
    add_user,
    get_connection
)

# Logging sozlamasi
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')

# --- BOT SOZLAMALARI ---
TOKEN = os.getenv("BOT_TOKEN")
if not TOKEN:
    raise ValueError("BOT_TOKEN topilmadi! Render Environment'ga qo'shing.")

bot = Bot(token=TOKEN)
dp = Dispatcher()
router = Router()

class FSM(StatesGroup):
    income_amount = State()
    expense_choice = State()

# --- BOT BUYRUQLAR MENYUSINI SOZLASH ---
async def set_bot_commands(bot: Bot):
    commands = [
        BotCommand(command="start", description="Botni ishga tushirish 🚀"),
        BotCommand(command="kirim", description="Balansga pul qo'shish 💰"),
        BotCommand(command="kirim_ochirish", description="Oxirgi kirimni o'chirish ❌"),
        BotCommand(command="balans", description="Joriy balanslarni tekshirish 💳"),
        BotCommand(command="balans_tozalash", description="Balansni noldan boshlash 🗑"),
        BotCommand(command="tozalash", description="Oxirgi xarajatni o'chirish 🗑"),
        BotCommand(command="kunlik", description="Kunlik hisobot 📊"),
        BotCommand(command="haftalik", description="Haftalik hisobot 📅"),
        BotCommand(command="oylik", description="Oylik hisobot 📈"),
        BotCommand(command="excel", description="Xarajatlarni Excel formatda yuklab olish 📊"),
        BotCommand(command="pdf", description="Xarajatlarni PDF formatda yuklab olish 📄")
    ]
    await bot.set_my_commands(commands)

# --- STANDARD KATEGORIYALARni BAZAGA QO'SHISH ---
def seed_default_categories():
    categories = ["Magazin", "Zapravka", "Apteka", "Stroy magazin", "Boshqa"]
    conn = get_connection()
    cursor = conn.cursor()
    for cat in categories:
        cursor.execute("INSERT OR IGNORE INTO categories (name) VALUES (?)", (cat,))
    conn.commit()
    conn.close()

# --- AVTOMATIK KATEGORIYAGA AJRATISH ---
def determine_category(name: str) -> str:
    name_lower = name.lower()
    if any(w in name_lower for w in ["benzin", "metan", "propan", "zapravka", "ai-92", "ai-95", "gaz"]):
        return "Zapravka"
    elif any(w in name_lower for w in ["dori", "tabletka", "apteka", "vitamin", "salfetka", "shpris"]):
        return "Apteka"
    elif any(w in name_lower for w in ["sement", "kraska", "mix", "truba", "kafel", "shurup", "bolt", "qum"]):
        return "Stroy magazin"
    elif any(w in name_lower for w in ["non", "shakar", "un", "kartoshka", "sariyog", "yog'", "yog", "sut", "choy", "go'sht", "gosht", "tuxum", "guruch", "makaron", "kolbasa", "sir", "tuz", "meva", "sabzi", "piyoz", "garox"]):
        return "Magazin"
    else:
        return "Boshqa"

# ================= 1. KIRIM VA BALANS QISMI =================

@router.message(Command("start"))
async def cmd_start(message: types.Message):
    user_id = message.from_user.id
    add_user(user_id)
    
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT card_balance, cash_balance, balance FROM users WHERE user_id = ?", (user_id,))
    row = cursor.fetchone()
    conn.close()
    
    card_bal, cash_bal, total_bal = row if row else (0, 0, 0)

    await message.answer(
        f"Assalomu alaykum! Xarajatlarni hisoblab boruvchi botga xush kelibsiz. 🚀\n\n"
        f"💳 <b>Plastik karta:</b> {card_bal:,.0f} so'm\n"
        f"💵 <b>Naqd pul:</b> {cash_bal:,.0f} so'm\n"
        f"💰 <b>Jami balans:</b> {total_bal:,.0f} so'm\n\n"
        f"📥 <b>Kirim qilish uchun:</b> <code>/kirim</code> buyrug'ini bosing\n"
        f"❌ <b>Oxirgi kirimni o'chirish:</b> <code>/kirim_ochirish</code>\n"
        f"🗑 <b>Balansni tozalash:</b> <code>/balans_tozalash</code>\n"
        f"📊 <b>Excel hisobot:</b> <code>/excel</code>\n"
        f"📄 <b>PDF hisobot:</b> <code>/pdf</code>\n\n"
        f"🛒 <b>Xarajat qilish:</b>\n"
        f"1️⃣ <code>non 2 ta 3500</code>\n"
        f"2️⃣ <code>sariyog 15000</code>",
        parse_mode="HTML"
    )

@router.message(Command("balans"))
async def cmd_balans(message: types.Message):
    user_id = message.from_user.id
    add_user(user_id)
    
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT card_balance, cash_balance, balance FROM users WHERE user_id = ?", (user_id,))
    row = cursor.fetchone()
    conn.close()
    
    card_bal, cash_bal, total_bal = row if row else (0, 0, 0)
    
    await message.answer(
        f"💳 <b>Balans hisoboti:</b>\n\n"
        f"💳 Plastik karta: {card_bal:,.0f} so'm\n"
        f"💵 Naqd pul: {cash_bal:,.0f} so'm\n"
        f"━━━━━━━━━━━━━━━━━━━\n"
        f"💰 <b>Jami balans: {total_bal:,.0f} so'm</b>",
        parse_mode="HTML"
    )

@router.message(Command("balans_tozalash"))
async def cmd_balans_tozalash(message: types.Message):
    user_id = message.from_user.id
    add_user(user_id)
    
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("UPDATE users SET card_balance = 0, cash_balance = 0, balance = 0, last_income = 0, last_income_type = 'cash' WHERE user_id = ?", (user_id,))
    conn.commit()
    conn.close()
    
    await message.answer("🗑 <b>Barcha balansingiz (Plastik va Naqd) 0 so'm qilib tozalandi!</b>", parse_mode="HTML")

@router.message(Command("kirim"))
async def cmd_kirim(message: types.Message):
    user_id = message.from_user.id
    add_user(user_id)
    
    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(text="💳 Plastik karta", callback_data="inc_card"),
            InlineKeyboardButton(text="💵 Naqd pul", callback_data="inc_cash")
        ]
    ])
    await message.answer("📥 Qaysi balansga pul qo'shmoqchisiz? Tanlang:", reply_markup=keyboard)

@router.callback_query(F.data.in_({"inc_card", "inc_cash"}))
async def process_income_choice(callback: types.CallbackQuery, state: FSMContext):
    inc_type = "card" if callback.data == "inc_card" else "cash"
    type_name = "Plastik karta" if inc_type == "card" else "Naqd pul"
    
    await state.update_data(income_type=inc_type)
    await state.set_state(FSM.income_amount)
    
    await callback.message.edit_text(f"💰 <b>{type_name}</b> uchun summani kiriting (masalan: 1000000):", parse_mode="HTML")
    await callback.answer()

@router.message(StateFilter(FSM.income_amount), F.text)
async def process_income_amount(message: types.Message, state: FSMContext):
    # Foydalanuvchi kiritgan raqamli xabarni avtomatik o'chiramiz
    try:
        await message.delete()
    except Exception:
        pass

    text = re.sub(r'\s+', '', message.text)

    if text.isdigit():
        amount = float(text)
        user_id = message.from_user.id
        data = await state.get_data()
        inc_type = data.get("income_type", "cash")
        
        add_user(user_id)
        conn = get_connection()
        cursor = conn.cursor()
        
        if inc_type == "card":
            cursor.execute("UPDATE users SET card_balance = card_balance + ?, last_income = ?, last_income_type = ? WHERE user_id = ?", (amount, amount, 'card', user_id))
        else:
            cursor.execute("UPDATE users SET cash_balance = cash_balance + ?, last_income = ?, last_income_type = ? WHERE user_id = ?", (amount, amount, 'cash', user_id))
            
        cursor.execute("UPDATE users SET balance = card_balance + cash_balance WHERE user_id = ?", (user_id,))
        conn.commit()
        
        cursor.execute("SELECT card_balance, cash_balance, balance FROM users WHERE user_id = ?", (user_id,))
        card_bal, cash_bal, total_bal = cursor.fetchone()
        conn.close()

        type_label = "Plastik karta" if inc_type == "card" else "Naqd pul"
        await message.answer(
            f"✅ <b>{type_label}</b>ga {amount:,.0f} so'm qo'shildi.\n\n"
            f"💳 Plastik: {card_bal:,.0f} so'm\n"
            f"💵 Naqd: {cash_bal:,.0f} so'm\n"
            f"💰 Jami balans: {total_bal:,.0f} so'm",
            parse_mode="HTML"
        )
        await state.clear()
    else:
        await message.answer("❌ Noto'g'ri summa kiritildi. Iltimos, faqat raqam kiriting:")

@router.message(Command("kirim_ochirish"))
async def cmd_kirim_ochirish(message: types.Message):
    user_id = message.from_user.id
    add_user(user_id)
    
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT card_balance, cash_balance, last_income, last_income_type FROM users WHERE user_id = ?", (user_id,))
    row = cursor.fetchone()
    
    if not row or row[2] <= 0:
        await message.answer("⚠️ O'chirish uchun oxirgi kirim topilmadi yoki allaqachon bekor qilingan.")
        conn.close()
        return
        
    card_bal, cash_bal, last_income, last_income_type = row
    
    if last_income_type == 'card':
        card_bal -= last_income
    else:
        cash_bal -= last_income
        
    cursor.execute(
        "UPDATE users SET card_balance = ?, cash_balance = ?, balance = ?, last_income = 0, last_income_type = 'cash' WHERE user_id = ?", 
        (card_bal, cash_bal, card_bal + cash_bal, user_id)
    )
    conn.commit()
    conn.close()
    
    type_label = "Plastik karta" if last_income_type == 'card' else "Naqd pul"
    await message.answer(
        f"❌ <b>Oxirgi kirim bekor qilindi ({type_label}):</b>\n"
        f"🔸 {last_income:,.0f} so'm ayrildi.\n\n"
        f"💳 Plastik: {card_bal:,.0f} so'm\n"
        f"💵 Naqd: {cash_bal:,.0f} so'm\n"
        f"💰 Jami balans: {card_bal + cash_bal:,.0f} so'm",
        parse_mode="HTML"
    )

# ================= 2. MENYU BUYRUQLARI, EXCEL VA PDF =================

@router.message(Command("tozalash"))
async def cmd_tozalash(message: types.Message):
    user_id = message.from_user.id
    
    conn = get_connection()
    cursor = conn.cursor()
    
    try:
        cursor.execute('''
            SELECT id, amount, item_name, payment_type FROM expenses 
            WHERE user_id = ? ORDER BY id DESC LIMIT 1
        ''', (user_id,))
        row = cursor.fetchone()
    except Exception:
        cursor.execute('''
            SELECT id, amount, item_name FROM expenses 
            WHERE user_id = ? ORDER BY id DESC LIMIT 1
        ''', (user_id,))
        row = cursor.fetchone()
        if row:
            row = (*row, 'cash')
            
    if row:
        exp_id, amount, item_name, pay_type = row
        cursor.execute('DELETE FROM expenses WHERE id = ?', (exp_id,))
        
        if pay_type == 'card':
            cursor.execute('UPDATE users SET card_balance = card_balance + ? WHERE user_id = ?', (amount, user_id))
        else:
            cursor.execute('UPDATE users SET cash_balance = cash_balance + ? WHERE user_id = ?', (amount, user_id))
            
        cursor.execute('UPDATE users SET balance = card_balance + cash_balance WHERE user_id = ?', (user_id,))
        conn.commit()
        
        cursor.execute("SELECT card_balance, cash_balance, balance FROM users WHERE user_id = ?", (user_id,))
        card_bal, cash_bal, total_bal = cursor.fetchone()
        
        pay_label = "💳 Plastik" if pay_type == 'card' else "💵 Naqd"
        await message.answer(
            f"🗑 <b>Oxirgi xarajat bekor qilindi:</b>\n"
            f"🔸 {item_name} — {int(amount):,} so'm\n"
            f"🔄 Summa <b>{pay_label}</b> hisobiga qaytarildi.\n\n"
            f"💳 Plastik: {card_bal:,.0f} so'm\n"
            f"💵 Naqd: {cash_bal:,.0f} so'm\n"
            f"💰 Jami balans: {total_bal:,.0f} so'm",
            parse_mode="HTML"
        )
    else:
        await message.answer("⚠️ Bekor qilish uchun xarajatlar tarixi topilmadi.")
    
    conn.close()

@router.message(Command("kunlik"))
async def cmd_kunlik(message: types.Message):
    user_id = message.from_user.id
    tz = pytz.timezone("Asia/Tashkent")
    today = datetime.now(tz).strftime("%Y-%m-%d")
    
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('''
        SELECT c.name, e.item_name, e.amount, e.time 
        FROM expenses e
        JOIN categories c ON e.category_id = c.id
        WHERE e.user_id = ? AND e.date = ?
    ''', (user_id, today))
    rows = cursor.fetchall()
    conn.close()
    
    if not rows:
        await message.answer(f"📅 <b>{today}</b>\n\nBugun uchun xarajatlar yo'q. 🤷‍♂️", parse_mode="HTML")
        return
        
    total = sum(row[2] for row in rows)
    grouped = {}
    for cat_name, item_name, amount, time in rows:
        if cat_name not in grouped:
            grouped[cat_name] = []
        grouped[cat_name].append((item_name, amount, time))
    
    report_lines = [f"📊 <b>Bugungi xarajatlar ({today}):</b>\n"]
    for cat, items in grouped.items():
        report_lines.append(f"📂 <b>{cat}:</b>")
        for item_name, amount, time in items:
            report_lines.append(f"  • {item_name} — {int(amount):,} so'm ({time})")
        report_lines.append("")
        
    report_lines.append(f"💰 <b>Jami kunlik xarajat: {int(total):,} so'm</b>")
    await message.answer("\n".join(report_lines), parse_mode="HTML")

@router.message(Command("haftalik"))
async def cmd_haftalik(message: types.Message):
    user_id = message.from_user.id
    tz = pytz.timezone("Asia/Tashkent")
    now = datetime.now(tz)
    today_str = now.strftime("%Y-%m-%d")
    start_of_week = (now - timedelta(days=now.weekday())).strftime("%Y-%m-%d")
    
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('''
        SELECT c.name, e.item_name, e.amount, e.date 
        FROM expenses e
        JOIN categories c ON e.category_id = c.id
        WHERE e.user_id = ? AND e.date BETWEEN ? AND ?
        ORDER BY e.date DESC
    ''', (user_id, start_of_week, today_str))
    rows = cursor.fetchall()
    conn.close()
    
    if not rows:
        await message.answer(f"📅 <b>Haftalik hisobot ({start_of_week} — {today_str})</b>\n\nXarajatlar mavjud emas. 🤷‍♂️", parse_mode="HTML")
        return
        
    total = sum(row[2] for row in rows)
    grouped = {}
    for cat_name, item_name, amount, date in rows:
        if cat_name not in grouped:
            grouped[cat_name] = []
        grouped[cat_name].append((item_name, amount, date))
    
    report_lines = [f"📅 <b>Haftalik xarajatlar ({start_of_week} — {today_str}):</b>\n"]
    for cat, items in grouped.items():
        report_lines.append(f"📂 <b>{cat}:</b>")
        for item_name, amount, date in items:
            report_lines.append(f"  • {item_name} — {int(amount):,} so'm ({date})")
        report_lines.append("")
        
    report_lines.append(f"💰 <b>Jami haftalik xarajat: {int(total):,} so'm</b>")
    await message.answer("\n".join(report_lines), parse_mode="HTML")

@router.message(Command("oylik"))
async def cmd_oylik(message: types.Message):
    user_id = message.from_user.id
    tz = pytz.timezone("Asia/Tashkent")
    now = datetime.now(tz)
    current_year_month = now.strftime("%Y-%m")
    
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('''
        SELECT c.name, e.item_name, e.amount, e.date 
        FROM expenses e
        JOIN categories c ON e.category_id = c.id
        WHERE e.user_id = ? AND e.date LIKE ?
        ORDER BY e.date DESC
    ''', (user_id, f"{current_year_month}%"))
    rows = cursor.fetchall()
    conn.close()
    
    if not rows:
        await message.answer(f"📈 <b>Oylik hisobot ({current_year_month})</b>\n\nXarajatlar mavjud emas. 🤷‍♂️", parse_mode="HTML")
        return
        
    total = sum(row[2] for row in rows)
    grouped = {}
    for cat_name, item_name, amount, date in rows:
        if cat_name not in grouped:
            grouped[cat_name] = []
        grouped[cat_name].append((item_name, amount, date))
    
    report_lines = [f"📈 <b>Oylik xarajatlar ({current_year_month}):</b>\n"]
    for cat, items in grouped.items():
        report_lines.append(f"📂 <b>{cat}:</b>")
        for item_name, amount, date in items:
            report_lines.append(f"  • {item_name} — {int(amount):,} so'm ({date})")
        report_lines.append("")
        
    report_lines.append(f"💰 <b>Jami oylik xarajat: {int(total):,} so'm</b>")
    await message.answer("\n".join(report_lines), parse_mode="HTML")

@router.message(Command("excel"))
async def cmd_excel_report(message: types.Message):
    user_id = message.from_user.id
    
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('''
        SELECT c.name, e.item_name, e.amount, e.date, e.time, e.payment_type 
        FROM expenses e
        JOIN categories c ON e.category_id = c.id
        WHERE e.user_id = ?
        ORDER BY e.date DESC, e.time DESC
    ''', (user_id,))
    rows = cursor.fetchall()
    conn.close()
    
    if not rows:
        await message.answer("⚠️ Excel hisobot yaratish uchun xarajatlar tarixi topilmadi.")
        return
        
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Xarajatlar tarixi"
    
    headers = ["Kategoriya", "Nomi / Tavsif", "Summa (so'm)", "To'lov turi", "Sana", "Vaqt"]
    ws.append(headers)
    
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        
    for row_data in rows:
        cat_name, item_name, amount, date_val, time_val, pay_type = row_data
        pay_label = "Plastik" if pay_type == "card" else "Naqd"
        ws.append([cat_name, item_name, amount, pay_label, date_val, time_val])
        
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=6):
        for cell in row:
            cell.border = thin_border
            if cell.column == 3:
                cell.number_format = '#,##0'
                cell.alignment = align_right
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    
    document = BufferedInputFile(file_stream.getvalue(), filename="xarajatlar_hisoboti.xlsx")
    
    await message.answer_document(
        document=document,
        caption="📊 Mana sizning barcha xarajatlaringiz jamlangan **Excel hisobot** faylingiz!",
        parse_mode="Markdown"
    )

@router.message(Command("pdf"))
async def cmd_pdf_report(message: types.Message):
    user_id = message.from_user.id
    
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('''
        SELECT c.name, e.item_name, e.amount, e.date, e.time, e.payment_type 
        FROM expenses e
        JOIN categories c ON e.category_id = c.id
        WHERE e.user_id = ?
        ORDER BY e.date DESC, e.time DESC
    ''', (user_id,))
    rows = cursor.fetchall()
    conn.close()
    
    if not rows:
        await message.answer("⚠️ PDF hisobot yaratish uchun xarajatlar tarixi topilmadi.")
        return
        
    file_stream = io.BytesIO()
    doc = SimpleDocTemplate(file_stream, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    story = []
    
    font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    if os.path.exists(font_path):
        pdfmetrics.registerFont(TTFont('DejaVu', font_path))
        font_name = 'DejaVu'
    else:
        font_name = 'Helvetica'
        
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading1'],
        fontName=font_name,
        fontSize=16,
        alignment=1,
        spaceAfter=20
    )
    
    cell_style = ParagraphStyle(
        'CellStyle',
        parent=styles['Normal'],
        fontName=font_name,
        fontSize=10
    )
    
    header_style = ParagraphStyle(
        'HeaderStyle',
        parent=styles['Normal'],
        fontName=font_name,
        fontSize=10,
        textColor=colors.whitesmoke,
        alignment=1
    )

    story.append(Paragraph("<b>Xarajatlar tarixi (Hisobot)</b>", title_style))
    story.append(Spacer(1, 10))
    
    table_data = [[
        Paragraph("<b>Kategoriya</b>", header_style),
        Paragraph("<b>Nomi</b>", header_style),
        Paragraph("<b>Summa</b>", header_style),
        Paragraph("<b>To'lov</b>", header_style),
        Paragraph("<b>Sana / Vaqt</b>", header_style)
    ]]
    
    total_sum = 0
    for row_data in rows:
        cat_name, item_name, amount, date_val, time_val, pay_type = row_data
        pay_label = "Plastik" if pay_type == "card" else "Naqd"
        total_sum += amount
        
        table_data.append([
            Paragraph(str(cat_name), cell_style),
            Paragraph(str(item_name), cell_style),
            Paragraph(f"{int(amount):,} so'm", cell_style),
            Paragraph(pay_label, cell_style),
            Paragraph(f"{date_val} {time_val}", cell_style)
        ])
        
    t = Table(table_data, colWidths=[100, 140, 90, 70, 100])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#4F81BD")),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9F9F9")]),
    ]))
    
    story.append(t)
    story.append(Spacer(1, 15))
    
    summary_style = ParagraphStyle(
        'SummaryStyle',
        parent=styles['Normal'],
        fontName=font_name,
        fontSize=12,
        alignment=2
    )
    story.append(Paragraph(f"<b>Jami xarajat: {int(total_sum):,} so'm</b>", summary_style))
    
    doc.build(story)
    file_stream.seek(0)
    
    document = BufferedInputFile(file_stream.getvalue(), filename="xarajatlar_hisoboti.pdf")
    
    await message.answer_document(
        document=document,
        caption="📄 Mana sizning barcha xarajatlaringiz jamlangan **PDF hisobot** faylingiz!",
        parse_mode="Markdown"
    )

# ================= 3. XARAJATLARNI MATNDAN O'QISH VA TANLOV =================

@router.message(StateFilter(None), F.text)
async def process_text_message(message: types.Message, state: FSMContext):
    if message.text.startswith('/'):
        return

    # Foydalanuvchi yozgan matnni avtomatik o'chirib tashlaymiz
    try:
        await message.delete()
    except Exception:
        pass

    user_id = message.from_user.id
    add_user(user_id)
    lines = message.text.strip().split('\n')

    pattern_unit = re.compile(r"^(.*?)\s+(\d+(?:\.\d+)?)\s*(ta|kg|l|litr|m|metr)\s+([\d\s]+)$", re.IGNORECASE)
    pattern_simple = re.compile(r"^(.*?)\s+([\d\s]+)$", re.IGNORECASE)

    parsed_expenses = []
    total_expense = 0
    
    tz = pytz.timezone("Asia/Tashkent")
    now = datetime.now(tz)
    date_str = now.strftime("%Y-%m-%d")
    time_str = now.strftime("%H:%M")

    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT id, name FROM categories")
        cat_rows = cursor.fetchall()
        conn.close()
        cat_dict = {name.lower(): cat_id for cat_id, name in cat_rows}

        for line in lines:
            line_text = line.strip()
            if not line_text:
                continue

            match_unit = pattern_unit.match(line_text)
            match_simple = pattern_simple.match(line_text)

            if match_unit:
                name = match_unit.group(1).strip().capitalize()
                qty_str = match_unit.group(2)
                qty = float(qty_str) if '.' in qty_str else int(qty_str)
                unit = match_unit.group(3).strip().lower()
                price_str = match_unit.group(4)
                price = float(re.sub(r'\s+', '', price_str))
                
                line_total = qty * price
                item_full_name = f"{name} {qty} {unit}"

                cat_name = determine_category(name)
                cat_id = cat_dict.get(cat_name.lower(), 1)

                parsed_expenses.append({
                    "category": cat_name, 
                    "cat_id": cat_id, 
                    "name": item_full_name, 
                    "amount": line_total,
                    "date_str": date_str,
                    "time_str": time_str
                })
                total_expense += line_total

            elif match_simple:
                name = match_simple.group(1).strip().capitalize()
                price_str = match_simple.group(2)
                line_total = float(re.sub(r'\s+', '', price_str))
                item_full_name = name

                cat_name = determine_category(name)
                cat_id = cat_dict.get(cat_name.lower(), 1)

                parsed_expenses.append({
                    "category": cat_name, 
                    "cat_id": cat_id, 
                    "name": item_full_name, 
                    "amount": line_total,
                    "date_str": date_str,
                    "time_str": time_str
                })
                total_expense += line_total

    except Exception as e:
        logging.error(f"Process text error: {e}", exc_info=True)
        await message.answer("❌ Xatolik yuz berdi. Iltimos, xabarni to'g'ri formatda yuboring.")
        return

    if not parsed_expenses:
        await message.answer(
            "⚠️ Xabarni to'g'ri formatda kiriting:\n\n"
            "🛒 Xarajat: <code>non 2 ta 3500</code> yoki <code>sariyog 15000</code>",
            parse_mode="HTML"
        )
        return

    await state.update_data(pending_expenses=parsed_expenses, total_expense=total_expense)
    await state.set_state(FSM.expense_choice)

    grouped = {}
    for exp in parsed_expenses:
        c = exp["category"]
        if c not in grouped:
            grouped[c] = []
        grouped[c].append(exp)

    preview_text = "🛒 <b>Xarajatlar ro'yxati:</b>\n"
    for cat, items in grouped.items():
        preview_text += f"📂 <b>{cat}:</b>\n"
        for item in items:
            preview_text += f"  • {item['name']} — {int(item['amount']):,} so'm\n"
        preview_text += "\n"

    preview_text += f"💰 <b>Jami xarajat: {int(total_expense):,} so'm</b>\n\n👇 <b>Ushbu xarajatni qaysi hisobdan ayiramiz?</b>"

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(text="💳 Plastik karta", callback_data="exp_card"),
            InlineKeyboardButton(text="💵 Naqd pul", callback_data="exp_cash")
        ],
        [
            InlineKeyboardButton(text="❌ Bekor qilish", callback_data="exp_cancel")
        ]
    ])

    await message.answer(preview_text, parse_mode="HTML", reply_markup=keyboard)

# ================= 4. XARAJATNI TASDIQLASH (CALLBACK) =================

@router.callback_query(StateFilter(FSM.expense_choice), F.data.in_({"exp_card", "exp_cash", "exp_cancel"}))
async def process_expense_choice(callback: types.CallbackQuery, state: FSMContext):
    if callback.data == "exp_cancel":
        await callback.message.edit_text("❌ <b>Xarajat yozish bekor qilindi.</b>", parse_mode="HTML")
        await state.clear()
        return
        
    data = await state.get_data()
    pending_expenses = data.get("pending_expenses", [])
    total_expense = data.get("total_expense", 0)
    user_id = callback.from_user.id
    
    pay_type = "card" if callback.data == "exp_card" else "cash"
    
    conn = get_connection()
    cursor = conn.cursor()
    
    for exp in pending_expenses:
        cursor.execute(
            "INSERT INTO expenses (user_id, amount, category_id, item_name, date, time, payment_type) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (user_id, exp['amount'], exp['cat_id'], exp['name'], exp['date_str'], exp['time_str'], pay_type)
        )
        
    if pay_type == "card":
        cursor.execute("UPDATE users SET card_balance = card_balance - ? WHERE user_id = ?", (total_expense, user_id))
    else:
        cursor.execute("UPDATE users SET cash_balance = cash_balance - ? WHERE user_id = ?", (total_expense, user_id))
        
    cursor.execute("UPDATE users SET balance = card_balance + cash_balance WHERE user_id = ?", (user_id,))
    conn.commit()
    
    cursor.execute("SELECT card_balance, cash_balance, balance FROM users WHERE user_id = ?", (user_id,))
    card_bal, cash_bal, total_bal = cursor.fetchone()
    conn.close()

    grouped = {}
    for exp in pending_expenses:
        c = exp["category"]
        if c not in grouped:
            grouped[c] = []
        grouped[c].append(exp)

    response_parts = ["✅ <b>Xarajatlar muvaffaqiyatli saqlandi:</b>\n"]
    for cat, items in grouped.items():
        response_parts.append(f"📂 <b>{cat}:</b>")
        for item in items:
            response_parts.append(f"  • {item['name']} — {int(item['amount']):,} so'm")
        response_parts.append("")
        
    pay_label = "💳 Plastik" if pay_type == "card" else "💵 Naqd"
    response_parts.append(f"💰 <b>Jami: {int(total_expense):,} so'm</b> ({pay_label}dan ayrildi)")
    
    response_parts.append(
        f"━━━━━━━━━━━━━━━━━━━\n"
        f"💳 Plastik: {card_bal:,.0f} so'm\n"
        f"💵 Naqd: {cash_bal:,.0f} so'm\n"
        f"💰 <b>Joriy balans: {total_bal:,.0f} so'm</b>"
    )
    
    await callback.message.edit_text("\n".join(response_parts), parse_mode="HTML")
    await state.clear()

# ================= ASOSIY ISHGA TUSHIRISH =================

async def main():
    init_db()
    seed_default_categories()

    conn = get_connection()
    cursor = conn.cursor()
    
    try: cursor.execute("ALTER TABLE users ADD COLUMN card_balance REAL DEFAULT 0")
    except Exception: pass
    
    try: cursor.execute("ALTER TABLE users ADD COLUMN cash_balance REAL DEFAULT 0")
    except Exception: pass
    
    try: cursor.execute("ALTER TABLE users ADD COLUMN last_income REAL DEFAULT 0")
    except Exception: pass
    
    try: cursor.execute("ALTER TABLE users ADD COLUMN last_income_type TEXT DEFAULT 'cash'")
    except Exception: pass

    try: cursor.execute("ALTER TABLE expenses ADD COLUMN payment_type TEXT DEFAULT 'cash'")
    except Exception: pass

    conn.commit()
    conn.close()

    dp.include_router(router)
    await bot.delete_webhook(drop_pending_updates=True)
    await set_bot_commands(bot)
    
    logging.info("Telegram bot ishga tushdi...")
    await dp.start_polling(bot)

if __name__ == "__main__":
    try:
        asyncio.run(main())
    except (KeyboardInterrupt, SystemExit):
        logging.info("Bot to'xtatildi.")
